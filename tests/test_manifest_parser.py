"""Tests for the transport-manifest extraction layer and its Excel output."""
from __future__ import annotations

import json

import pytest
from PIL import Image

from services.excel_export import LOW_CONFIDENCE_THRESHOLD, build_manifest_workbook
from services.manifest_parser import (
    MANIFEST_PROVIDER_REQUIRED,
    ManifestParser,
    ManifestRow,
    ParsedManifest,
    ParseError,
    to_number_or_blank,
)
from services.ocr_service import OCRProvider, OCRResponse

SHL_MANIFEST = {
    "manifest_no": "2264",
    "manifest_date": "2026-10-06",
    "from_station": "K. Gate, Delhi",
    "to_station": "Raipur (MP)",
    "rows": [
        {
            "s_no": 1, "c_no": "148504", "station": "Bhilai", "consignee": "Soni Auto",
            "weight": 3, "no_of_pkg": 1, "to_pay": "", "paid": "Paid", "amount": "",
            "confidence": 0.9,
        },
        {
            "s_no": 2, "c_no": "148503", "station": "Raipur", "consignee": "Brijlal Mahajan",
            "weight": 2, "no_of_pkg": 1, "to_pay": 240, "paid": "", "amount": "",
            "confidence": 0.55,
        },
    ],
    "overall_confidence": 0.8,
}


class FakeStructuredProvider(OCRProvider):
    name = "fake"
    supports_structured = True

    def __init__(self, config=None, payload=None):
        super().__init__(config or {})
        self.payload = payload if payload is not None else SHL_MANIFEST

    def extract_text(self, images):
        raise NotImplementedError

    def extract_structured(self, images, prompt, schema_hint):
        return OCRResponse(provider=self.name, model="fake-1", structured=self.payload)


class FakeTextOnlyProvider(OCRProvider):
    name = "fake-text"
    supports_structured = False

    def extract_text(self, images):
        return OCRResponse(provider=self.name, text="")

    def extract_structured(self, images, prompt, schema_hint):
        raise NotImplementedError


def one_blank_image():
    return [Image.new("RGB", (10, 10), "white")]


class TestNumberCoercion:
    @pytest.mark.parametrize(
        "raw,expected",
        [
            (240, 240), ("240", 240), ("1,440", 1440), ("", ""),
            ("Paid", "Paid"), ("paid", "Paid"), ("X", ""), ("x", ""),
            ("-", ""), (None, ""), (5.5, 5.5), ("5.0", 5),
        ],
    )
    def test_to_number_or_blank(self, raw, expected):
        assert to_number_or_blank(raw) == expected


class TestManifestParser:
    def test_requires_a_structured_provider(self):
        with pytest.raises(ParseError, match="vision engine"):
            ManifestParser(FakeTextOnlyProvider())

    def test_error_message_names_the_fix(self):
        assert "Settings" in MANIFEST_PROVIDER_REQUIRED
        assert "Auto" in MANIFEST_PROVIDER_REQUIRED

    def test_parses_header_and_rows(self):
        parser = ManifestParser(FakeStructuredProvider())
        result = parser.parse_images(one_blank_image())
        assert result.manifest_no == "2264"
        assert result.from_station == "K. Gate, Delhi"
        assert result.to_station == "Raipur (MP)"
        assert len(result.rows) == 2

    def test_paid_mark_becomes_the_word_paid(self):
        parser = ManifestParser(FakeStructuredProvider())
        result = parser.parse_images(one_blank_image())
        assert result.rows[0].paid == "Paid"
        assert result.rows[0].to_pay == ""

    def test_numeric_to_pay_stays_numeric(self):
        parser = ManifestParser(FakeStructuredProvider())
        result = parser.parse_images(one_blank_image())
        assert result.rows[1].to_pay == 240

    def test_no_images_raises(self):
        parser = ManifestParser(FakeStructuredProvider())
        with pytest.raises(ParseError, match="No page images"):
            parser.parse_images([])

    def test_rows_outside_the_list_are_ignored(self):
        payload = dict(SHL_MANIFEST)
        payload["rows"] = [*SHL_MANIFEST["rows"], "not a row", None, 42]
        parser = ManifestParser(FakeStructuredProvider(payload=payload))
        result = parser.parse_images(one_blank_image())
        assert len(result.rows) == 2

    def test_trailing_comma_in_a_text_response_is_tolerated(self):
        class TextOnlyJsonProvider(OCRProvider):
            name = "textjson"
            supports_structured = True

            def extract_text(self, images):
                raise NotImplementedError

            def extract_structured(self, images, prompt, schema_hint):
                # No `structured` field (as a provider that only returns raw
                # text would give us) and a trailing comma before the closing
                # brace - local_json_repair must strip that before parsing.
                # (manifest_no ends in a letter, not a digit, to steer clear
                # of local_json_repair's unrelated "strip a stray quote after
                # a number" rule, which a trailing digit would also trip.)
                return OCRResponse(
                    provider=self.name, model="m",
                    raw_response='{"manifest_no": "PENDING", "rows": [],}',
                )

        parser = ManifestParser(TextOnlyJsonProvider())
        result = parser.parse_images(one_blank_image())
        assert result.manifest_no == "PENDING"
        assert result.rows == []

    def test_a_transient_provider_error_is_retried(self):
        class FlakyThenFixedProvider(OCRProvider):
            name = "flaky"
            supports_structured = True
            calls = 0

            def extract_text(self, images):
                raise NotImplementedError

            def extract_structured(self, images, prompt, schema_hint):
                self.calls += 1
                if self.calls == 1:
                    from services.ocr_service import OCRError
                    raise OCRError("timed out")
                return OCRResponse(provider=self.name, model="m", structured=SHL_MANIFEST)

        provider = FlakyThenFixedProvider()
        parser = ManifestParser(provider)
        result = parser.parse_images(one_blank_image())
        assert result.manifest_no == "2264"
        assert provider.calls == 2
        assert any("timed out" in w for w in result.warnings)


class TestManifestWorkbook:
    def _parsed(self):
        parser = ManifestParser(FakeStructuredProvider())
        return parser.parse_images(one_blank_image())

    def test_builds_a_workbook_with_one_row_per_consignment(self, tmp_path):
        import openpyxl

        path = build_manifest_workbook(self._parsed(), destination=tmp_path)
        wb = openpyxl.load_workbook(path)
        sheet = wb.active
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        assert "148504" in values
        assert "148503" in values
        assert "Soni Auto" in values

    def test_low_confidence_rows_are_highlighted(self, tmp_path):
        import openpyxl

        parsed = self._parsed()
        assert parsed.rows[1].confidence < LOW_CONFIDENCE_THRESHOLD
        path = build_manifest_workbook(parsed, destination=tmp_path)
        wb = openpyxl.load_workbook(path)
        sheet = wb.active
        # Find the row holding the low-confidence consignee and check its fill.
        target_row = None
        for row in sheet.iter_rows():
            if any(cell.value == "Brijlal Mahajan" for cell in row):
                target_row = row
                break
        assert target_row is not None
        assert target_row[0].fill.fgColor.rgb in ("00FFF2CC", "FFFFF2CC")

    def test_empty_manifest_still_builds(self, tmp_path):
        import openpyxl

        empty = ParsedManifest(manifest_no="9999", rows=[])
        path = build_manifest_workbook(empty, destination=tmp_path)
        wb = openpyxl.load_workbook(path)
        assert wb.active is not None


# --------------------------------------------------------------------------
# Several bills, one workbook
# --------------------------------------------------------------------------
def _manifest(no, rows, date="2026-10-07", frm="Kolkata", to="Raipur"):
    return ParsedManifest(
        manifest_no=no, manifest_date=date, from_station=frm, to_station=to,
        rows=[ManifestRow(**r) for r in rows], overall_confidence=0.8,
    )


def _bill(no, weights, to_pays, conf=0.9):
    return _manifest(no, [
        dict(s_no=i + 1, c_no=f"{no}{i}", station="Raipur", consignee=f"Party {i}",
             weight=w, no_of_pkg=1, to_pay=t, confidence=conf)
        for i, (w, t) in enumerate(zip(weights, to_pays))
    ])


class TestRowsFromRecords:
    def test_strings_from_the_editor_become_numbers_again(self):
        from services.manifest_parser import rows_from_records

        rows = rows_from_records([{
            "S.No": "1", "C.No.": "148504", "Station": "Raipur", "Consignee": "X",
            "Weight": "12.5", "No. of Pkg": "4", "To Pay": "240", "Paid": "Paid",
            "Amount": "", "Confidence": 55,
        }])
        assert rows[0].weight == 12.5 and rows[0].no_of_pkg == 4
        assert rows[0].to_pay == 240 and isinstance(rows[0].to_pay, int)
        assert rows[0].paid == "Paid" and rows[0].amount == ""
        assert rows[0].confidence == pytest.approx(0.55)

    def test_nan_and_none_become_blank_not_the_word_nan(self):
        from services.manifest_parser import rows_from_records

        rows = rows_from_records([{
            "S.No": 1, "C.No.": float("nan"), "Station": None, "Consignee": "Y",
            "Weight": float("nan"), "To Pay": None, "Confidence": float("nan"),
        }])
        assert rows[0].c_no == "" and rows[0].station == ""
        assert rows[0].weight == "" and rows[0].to_pay == ""
        assert rows[0].confidence == 0.0

    def test_a_completely_empty_added_row_is_dropped(self):
        from services.manifest_parser import rows_from_records

        rows = rows_from_records([
            {"S.No": None, "C.No.": None, "Station": None, "Consignee": None},
            {"S.No": 2, "C.No.": "9", "Station": "Durg", "Consignee": "Z"},
        ])
        assert len(rows) == 1 and rows[0].station == "Durg"


class TestMultiManifestWorkbook:
    def test_one_bill_is_the_same_single_sheet_as_before(self, tmp_path):
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        path = build_multi_manifest_workbook([_bill("2264", [3, 2], [240, ""])], tmp_path)
        wb = load_workbook(path)
        assert wb.sheetnames == ["Manifest"]

    def test_several_bills_land_in_one_file(self, tmp_path):
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        bills = [_bill("2264", [3, 2], [240, ""]), _bill("6037", [10], [500]),
                 _bill("7001", [1, 1, 1], ["", "", 90])]
        path = build_multi_manifest_workbook(bills, tmp_path)
        assert len(list(tmp_path.glob("*.xlsx"))) == 1          # ONE file
        wb = load_workbook(path)
        assert wb.sheetnames[:2] == ["All Manifests", "Summary"]
        assert wb.sheetnames[2:] == ["1 - 2264", "2 - 6037", "3 - 7001"]

        sheet = wb["All Manifests"]
        header_row = next(r for r in sheet.iter_rows() if r[0].value == "Bill #")
        headers = [c.value for c in header_row]
        assert headers[:5] == ["Bill #", "Manifest No.", "Date", "From", "To"]
        data = [[c.value for c in r] for r in sheet.iter_rows(min_row=header_row[0].row + 1)
                if isinstance(r[0].value, int)]
        assert len(data) == 6                                   # 2 + 1 + 3 consignments
        assert [d[0] for d in data] == [1, 1, 2, 3, 3, 3]
        assert [d[1] for d in data] == ["2264", "2264", "6037", "7001", "7001", "7001"]

    def test_totals_and_summary_are_correct_after_recalculation(self, tmp_path):
        """Real formulas, so recalculate with LibreOffice and read the values."""
        import shutil, subprocess
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        if not shutil.which("soffice"):
            pytest.skip("LibreOffice not available to recalculate formulas")
        bills = [_bill("2264", [3, 2], [240, ""]), _bill("6037", [10], [500]),
                 _bill("7001", [1, 1, 1], ["", "", 90])]
        path = build_multi_manifest_workbook(bills, tmp_path)
        out = tmp_path / "recalc"
        out.mkdir()
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir",
                        str(out), path], check=True, capture_output=True, timeout=120)
        wb = load_workbook(out / path.split("/")[-1], data_only=True)

        total = next(r for r in wb["All Manifests"].iter_rows()
                     if str(r[0].value or "").startswith("GRAND TOTAL"))
        values = [c.value for c in total]
        assert 18 in values and 830 in values                   # weight 3+2+10+1+1+1, to-pay 240+500+90

        summary = [[c.value for c in r] for r in wb["Summary"].iter_rows()
                   if isinstance(r[0].value, int)]
        assert [(s[0], s[5], s[6], s[7]) for s in summary] == [
            (1, 2, 5, 240), (2, 1, 10, 500), (3, 3, 3, 90)]
        grand = next(r for r in wb["Summary"].iter_rows() if r[0].value == "TOTAL")
        assert (grand[5].value, grand[6].value, grand[7].value) == (6, 18, 830)

        for sheet in wb:                                        # no #REF!/#NAME? anywhere
            for row in sheet.iter_rows():
                for cell in row:
                    assert not (isinstance(cell.value, str) and cell.value.startswith("#")), (
                        sheet.title, cell.coordinate, cell.value)

    def test_same_manifest_number_twice_and_odd_names_get_safe_unique_sheets(self, tmp_path):
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        bills = [_bill("2264", [1], [1]), _bill("2264", [2], [2]),
                 _bill("A/B:C*D?[E]" + "x" * 40, [3], [3])]
        wb = load_workbook(build_multi_manifest_workbook(bills, tmp_path))
        names = wb.sheetnames[2:]
        assert len(set(n.lower() for n in names)) == 3
        assert all(len(n) <= 31 for n in names)
        assert not any(ch in n for n in names for ch in "[]:*?/\\")

    def test_low_confidence_rows_are_still_highlighted_in_the_combined_sheet(self, tmp_path):
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        bills = [_bill("1", [1], [1], conf=0.95), _bill("2", [2], [2], conf=0.4)]
        wb = load_workbook(build_multi_manifest_workbook(bills, tmp_path))
        sheet = wb["All Manifests"]
        rows = [r for r in sheet.iter_rows() if isinstance(r[0].value, int)]
        fills = [r[0].fill.fgColor.rgb for r in rows]
        assert fills[0] not in ("00FFF2CC", "FFFFF2CC")
        assert fills[1] in ("00FFF2CC", "FFFFF2CC")

    def test_nothing_to_export_is_an_error(self, tmp_path):
        from services.excel_export import build_multi_manifest_workbook

        with pytest.raises(ValueError):
            build_multi_manifest_workbook([], tmp_path)


class TestSerialNumbersContinueAcrossBills:
    def test_combined_sheet_numbers_run_on_from_bill_to_bill(self, tmp_path):
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        bills = [_bill("6059", [1] * 6, [""] * 6), _bill("6060", [1] * 6, [""] * 6),
                 _bill("6061", [1] * 3, [""] * 3)]
        wb = load_workbook(build_multi_manifest_workbook(bills, tmp_path))
        sheet = wb["All Manifests"]
        serials = [r[5].value for r in sheet.iter_rows() if isinstance(r[0].value, int)]
        assert serials == list(range(1, 16))                    # 6 + 6 + 3, no restart at 1

    def test_each_bills_own_sheet_keeps_the_number_printed_on_its_paper(self, tmp_path):
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        wb = load_workbook(build_multi_manifest_workbook(
            [_bill("1", [1, 1], ["", ""]), _bill("2", [1, 1], ["", ""])], tmp_path))
        own = [r[0].value for r in wb["2 - 2"].iter_rows() if isinstance(r[0].value, int)]
        assert own == [1, 2]


class TestHeaderNormalising:
    @pytest.mark.parametrize("raw, expected", [
        ("KOLKATA - 68, COTTON STREET", "KOLKATA"),
        ("From - KOLKATA - 68, COTTON STREET", "KOLKATA"),
        ("KOLKATA", "KOLKATA"),
        ("Raipur (MP)", "Raipur (MP)"),
        ("K. Gate, Delhi", "K. Gate, Delhi"),              # a comma alone is not an address
        ("", ""), (None, ""),
    ])
    def test_station_is_just_the_town(self, raw, expected):
        from services.manifest_parser import normalize_station

        assert normalize_station(raw) == expected

    @pytest.mark.parametrize("raw, expected", [
        ("05.10.26", "05/10/2026"), ("5/10/26", "05/10/2026"), ("05-10-2026", "05/10/2026"),
        ("2026-10-05", "05/10/2026"), ("05/10/2026", "05/10/2026"),
        ("garbage", "garbage"), ("31.02.26", "31.02.26"), ("", ""),
    ])
    def test_date_is_shown_as_dd_mm_yyyy_with_a_20yy_year(self, raw, expected):
        from services.manifest_parser import normalize_date

        assert normalize_date(raw) == expected

    def test_dates_in_the_workbook_are_real_dates_shown_dd_mm_yyyy(self, tmp_path):
        import datetime as dt
        from openpyxl import load_workbook
        from services.excel_export import build_multi_manifest_workbook

        bills = [_manifest("1", [dict(s_no=1, c_no="a", weight=1)], date="05.10.26"),
                 _manifest("2", [dict(s_no=1, c_no="b", weight=1)], date="2026-10-06"),
                 _manifest("3", [dict(s_no=1, c_no="c", weight=1)], date="not a date")]
        wb = load_workbook(build_multi_manifest_workbook(bills, tmp_path))
        cells = [r[2] for r in wb["All Manifests"].iter_rows() if isinstance(r[0].value, int)]
        assert cells[0].value == dt.datetime(2026, 10, 5) and cells[0].number_format == "DD/MM/YYYY"
        assert cells[1].value == dt.datetime(2026, 10, 6)
        assert cells[2].value == "not a date"                    # unreadable: shown as written
        assert wb["2 - 2"]["A2"].value == "Dated 06/10/2026 · Kolkata → Raipur"

    def test_the_same_station_reads_the_same_on_every_bill(self):
        payloads = [
            dict(SHL_MANIFEST, from_station="KOLKATA"),
            dict(SHL_MANIFEST, from_station="KOLKATA - 68, COTTON STREET"),
        ]
        froms = [ManifestParser(FakeStructuredProvider(payload=p)).parse_images(
            [Image.new("RGB", (4, 4))]).from_station for p in payloads]
        assert froms == ["KOLKATA", "KOLKATA"]
